from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_resume_text(file_bytes: bytes, filename: str, content_type: str | None = None) -> str:
    lower_filename = filename.lower()
    if lower_filename.endswith(".pdf") or content_type == "application/pdf":
        return _extract_pdf_text(file_bytes)
    if lower_filename.endswith(".docx") or content_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    }:
        return _extract_docx_text(file_bytes)
    raise ValueError("Only PDF and DOCX resumes are supported")


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()
